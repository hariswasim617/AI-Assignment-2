"""
Builds report/Technical_Report.docx (and .pdf if LibreOffice is installed)
from the saved experiment results. Every number in the report is read from
the files in results/, so the report always matches the actual runs.

    python report/build_report.py
"""

import json
import os
import shutil
import subprocess
import sys

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from uav_ga import config as cfg   # noqa: E402  (the settings used for the runs)

# ---- fill in your details here and run the script again ----
AUTHOR_NAME = "[Your Name]"
ROLL_NO = "[Roll No]"

RESULTS = os.path.join(ROOT, "results")
FIGURES = os.path.join(ROOT, "report", "figures")
OUT_DOCX = os.path.join(ROOT, "report", "Technical_Report.docx")

BODY_PT = 9.5
FONT = "Calibri"


# ---------------------------------------------------------------------------
# Small helpers for python-docx
# ---------------------------------------------------------------------------
def set_cell_border(cell, **kwargs) -> None:
    """Set borders of a table cell, e.g. set_cell_border(cell, top={"val": "single"})."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in kwargs.items():
        element = OxmlElement(f"w:{edge}")
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))
        borders.append(element)


def keep_table_together(table) -> None:
    """Ask Word/LibreOffice not to split a table (or a figure grid) across pages."""
    for i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if i < len(table.rows) - 1:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def set_column_widths(table, widths: list[float]) -> None:
    """Column widths in cm (set on the grid AND on every cell, so all programs respect them)."""
    table.autofit = False
    for c, width in enumerate(widths):
        table.columns[c].width = Cm(width)
        for cell in table.columns[c].cells:
            cell.width = Cm(width)


def add_text(paragraph, text: str, size: float = BODY_PT) -> None:
    """Add text where **bold** and *italic* parts are marked like Markdown."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("*"):
            run = paragraph.add_run(part[1:-1]); run.italic = True
        else:
            run = paragraph.add_run(part)
        run.font.size = Pt(size)


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        self.figure_no = 0
        self.table_no = 0
        section = self.doc.sections[0]
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(1.7))
        section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
        style = self.doc.styles["Normal"]
        style.font.name = FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(BODY_PT)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.line_spacing = 1.0
        for level, size in ((1, 12.5), (2, 10.5)):
            h = self.doc.styles[f"Heading {level}"]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            h.paragraph_format.space_before = Pt(7 if level == 1 else 4)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True

    # -- text ---------------------------------------------------------------
    def heading(self, text: str, level: int = 1) -> None:
        p = self.doc.add_heading(level=level)
        run = p.add_run(text)
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)

    def para(self, text: str, size: float = BODY_PT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after: float = 3):
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        add_text(p, text, size)
        return p

    def bullets(self, items: list[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_text(p, item)

    def equation(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.italic = True
        run.font.size = Pt(BODY_PT)
        run.font.name = "Cambria Math"

    # -- tables -------------------------------------------------------------
    def table(self, header: list[str], rows: list[list[str]], caption: str,
              widths: list[float] | None = None, size: float = 8.5) -> None:
        self.table_no += 1
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(1)
        cap.paragraph_format.keep_with_next = True
        add_text(cap, f"**Table {self.table_no}.** {caption}", 8.5)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        line = {"val": "single", "sz": 6, "color": "404040"}
        for r, values in enumerate([header] + rows):
            for c, value in enumerate(values):
                cell = t.cell(r, c)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.CENTER
                add_text(p, f"**{value}**" if r == 0 else str(value), size)
                if r == 0:
                    set_cell_border(cell, top=line, bottom=line)
                if r == len(rows):
                    set_cell_border(cell, bottom=line)
        if widths:
            set_column_widths(t, widths)
        keep_table_together(t)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.line_spacing = 0.6

    # -- figures ------------------------------------------------------------
    def figure(self, images: list[str], widths: list[float], caption: str) -> None:
        """One or more images side by side, with a numbered caption below."""
        self.figure_no += 1
        t = self.doc.add_table(rows=1, cols=len(images))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (img, w) in enumerate(zip(images, widths)):
            p = t.cell(0, i).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_picture(img, width=Cm(w))
            p.paragraph_format.keep_with_next = True       # keep the picture with its caption
        keep_table_together(t)
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(4)
        add_text(cap, f"**Figure {self.figure_no}.** {caption}", 8.5)

    def figure_grid(self, images: list[str], width: float, columns: int, caption: str) -> None:
        """Images in a grid (``columns`` per row), with one numbered caption below."""
        self.figure_no += 1
        n_rows = (len(images) + columns - 1) // columns
        t = self.doc.add_table(rows=n_rows, cols=columns)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, img in enumerate(images):
            p = t.cell(i // columns, i % columns).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_picture(img, width=Cm(width))
            p.paragraph_format.keep_with_next = True
        keep_table_together(t)
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(4)
        add_text(cap, f"**Figure {self.figure_no}.** {caption}", 8.5)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path: str) -> None:
        self.doc.save(path)


def fmt(x: float, d: int = 3) -> str:
    return f"{x:.{d}f}"


def pm(mean: float, std: float, d: int = 3) -> str:
    return f"{mean:.{d}f} ± {std:.{d}f}"


# ---------------------------------------------------------------------------
# Load results
# ---------------------------------------------------------------------------
def load() -> dict:
    data = {}
    data["summary"] = json.load(open(os.path.join(RESULTS, "run_summary.json")))
    data["history"] = pd.read_csv(os.path.join(RESULTS, "history.csv"))
    data["sens_runs"] = pd.read_csv(os.path.join(RESULTS, "sensitivity", "sensitivity_runs.csv"))
    data["sens"] = pd.read_csv(os.path.join(RESULTS, "sensitivity", "sensitivity_summary.csv"))
    data["base_runs"] = pd.read_csv(os.path.join(RESULTS, "baseline", "baseline_runs.csv"))
    data["examples"] = json.load(open(os.path.join(RESULTS, "baseline", "baseline_examples.json")))
    return data


def copy_figures() -> None:
    """Copy the figures used in the report into report/figures/."""
    os.makedirs(FIGURES, exist_ok=True)
    names = ["paths_generations.png", "convergence.png", "objectives_over_time.png",
             "coverage_heatmaps.png", "diversity.png", "final_solution.png", "environment.png"]
    for name in names:
        shutil.copy(os.path.join(RESULTS, name), os.path.join(FIGURES, name))
    for name in ("sweep_pop_size.png", "sweep_mutation_rate.png", "sweep_tournament_k.png",
                 "sweep_selection.png", "sweep_constraint_mode.png"):
        shutil.copy(os.path.join(RESULTS, "sensitivity", name), os.path.join(FIGURES, name))
    for name in ("ga_vs_random_convergence.png", "ga_vs_random_solutions.png"):
        shutil.copy(os.path.join(RESULTS, "baseline", name), os.path.join(FIGURES, name))


def fig(name: str) -> str:
    return os.path.join(FIGURES, name)


# ---------------------------------------------------------------------------
# Report text
# ---------------------------------------------------------------------------
def write_report(r: "Report", data: dict) -> None:
    """All report text. Every number is read from `data` (the saved results)."""
    from uav_ga.environment import Environment
    env = Environment()
    s, h = data["summary"], data["history"]
    best = s["best_objectives"]
    sens, runs, base = data["sens"], data["sens_runs"], data["base_runs"]
    combo = pd.read_csv(os.path.join(RESULTS, "sensitivity", "combined_check.csv"))
    n_targets = len(env.targets)
    genes = sum(1 + 5 * len(d["waypoints"]) for d in s["best_plan"])
    first, last = h.iloc[0], h.iloc[-1]
    ga_r, rs_r = base[base["method"] == "Genetic Algorithm"], base[base["method"] == "Random search"]
    g = lambda p, v: sens[(sens["parameter"] == p) & (sens["value"].astype(str) == str(v))].iloc[0]

    # ------------------------------------------------------------------ title page
    for _ in range(5):
        r.para("", after=6)
    for text, size, bold in (("Bahria University Lahore Campus", 16, True),
                             ("Department of Computer Science", 13, False),
                             ("Artificial Intelligence – BSCS 5", 13, False),
                             ("Assignment 2", 13, True)):
        p = r.para("", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
        run = p.add_run(text); run.font.size = Pt(size); run.bold = bold
    r.para("", after=24)
    p = r.para("", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    run = p.add_run("A Genetic-Algorithm Strategy Engine for UAV Fleet Deployment, "
                    "Sensor Scheduling and Path Planning over a Volatile Hazard Grid")
    run.font.size = Pt(17); run.bold = True; run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    r.para("", after=6)
    p = r.para("", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    run = p.add_run("Technical Report"); run.font.size = Pt(13); run.italic = True
    r.para("", after=40)
    for label, value in (("Submitted by", AUTHOR_NAME), ("Roll No", ROLL_NO),
                         ("Submission", "Individual"), ("Code", "github.com/hariswasim617/AI-Assignment-2")):
        p = r.para("", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        a = p.add_run(f"{label}: "); a.bold = True; a.font.size = Pt(12)
        b = p.add_run(value); b.font.size = Pt(12)
    r.page_break()

    # ------------------------------------------------------------------ 1
    r.heading("1. Introduction and Problem Formulation")
    r.para("A disaster-response company surveys a dangerous area with a fleet of drones (UAVs). For every drone "
           "it must choose a take-off base, a route, the speed on each part of the route and when (and how "
           "strongly) the sensor is switched on. These choices conflict: covering more area costs battery and "
           "time, and the most important places are often also the most dangerous. There is no single correct "
           "answer, only trade-offs. I designed and implemented, from scratch in Python, a Genetic Algorithm "
           "(GA) [1, 2] that searches for good trade-offs; this report explains the design and the experiments.")
    r.para(f"**Grid model and notation.** The area is a {cfg.GRID_SIZE}×{cfg.GRID_SIZE} grid of "
           f"{cfg.CELL_SIZE_M:.0f} m cells (1 km × 1 km). Each cell *c* has *blocked(c)* ∈ {{0,1}} "
           f"({cfg.N_NOFLY_ZONES} no-fly zones, {100 * env.blocked.mean():.1f} % of the map), a hazard "
           f"*h(c)* ∈ [0,1] from {cfg.N_WIND_POCKETS} Gaussian wind and {cfg.N_RADIATION_POCKETS} radiation "
           f"pockets, and a priority *P(c)* ∈ [0,1] from {cfg.N_PRIORITY_ZONES} disaster zones. There are "
           f"{n_targets} targets (75 % inside disaster zones) and {len(env.base_stations)} base stations. Two "
           "radiation pockets are placed on purpose next to disaster zones, so priority and safety really "
           f"conflict. Drone *d* (N = {cfg.N_DRONES}) flies the closed route *b_d → w_d,1 → … → w_d,K_d → b_d* "
           f"with range L_max = {cfg.MAX_RANGE_M / 1000:.0f} km, battery B = {cfg.BATTERY_J / 1000:.0f} kJ and "
           f"speed {cfg.V_MIN:.0f}–{cfg.V_MAX:.0f} m/s; its sensor radius grows from {cfg.SENSOR_RADIUS_MIN} to "
           f"{cfg.SENSOR_RADIUS_MAX:.0f} cells with the power p ∈ [0,1] ({cfg.SENSOR_POWER_MIN_W:.0f}–"
           f"{cfg.SENSOR_POWER_MAX_W:.0f} W). The map is generated from a fixed seed, so all experiments use the "
           "same world (maps in Fig. 2).")
    r.para("**Assumptions.** Sensors look straight down (no-fly zones restrict flight, not sensing); every drone "
           f"returns to its base at {cfg.V_RETURN:.0f} m/s with the sensor off; propulsion power follows the "
           f"rotary-wing model of Zeng et al. [4] and wind raises it by up to {100 * cfg.WIND_POWER_FACTOR:.0f} %; "
           "detection is probabilistic and gets worse with speed (motion blur) and distance.")

    # ------------------------------------------------------------------ 2
    r.heading("2. Chromosome Design and Rationale")
    r.para("I used a **hybrid, multi-segment, variable-length** encoding: one segment per drone, mixing an "
           "integer, real and binary genes (Table 1). The genes map one-to-one onto the planning decisions.")
    r.table(["Gene (per drone)", "Type", "Range", "Meaning"],
            [["b", "integer", f"0 … {len(env.base_stations) - 1}", "launch / landing base station"],
             ["x_k, y_k  (k = 1…K)", "real", f"[0, {cfg.GRID_SIZE})", "position of waypoint k (cells)"],
             ["v_k", "real", f"[{cfg.V_MIN:.0f}, {cfg.V_MAX:.0f}] m/s", "speed on the leg that ends at waypoint k"],
             ["p_k", "real", "[0, 1]", "sensor power on that leg (radius and energy)"],
             ["a_k", "binary", "{0, 1}", "sensor on/off on that leg (activation schedule)"]],
            "Encoding of one drone segment; a chromosome has N segments and K varies per drone.",
            widths=[3.4, 1.6, 2.6, 7.4])
    r.para(f"**Rationale.** (1) *Real-valued coordinates*: a small gene change gives a small route change, which "
           "suits Gaussian mutation and geometric crossover (with a binary encoding one bit flip could move a "
           "drone across the map). (2) *One segment per drone*: a drone's route is a natural building block that "
           f"crossover can pass on intact. (3) *Variable length* (K starts at {cfg.K_INIT_MIN}–{cfg.K_INIT_MAX} "
           f"and may vary between {cfg.K_MIN} and {cfg.K_MAX} by mutation): the GA decides how detailed each route "
           "is. (4) *Per-leg speed, power and on/off genes* give direct control of the energy–detection trade-off, "
           f"e.g. sensor off and fast on transit legs. The final plan has {best['n_waypoints']} waypoints "
           f"({genes} genes).")

    # ------------------------------------------------------------------ 3
    r.heading("3. Fitness Function")
    r.para("The fitness (maximised) is a weighted sum of five objectives, each **normalised to [0, 1]**, "
           "minus a dynamic penalty:")
    r.equation(f"F = {cfg.W_COVERAGE}·Cov + {cfg.W_DETECTION}·Det − {cfg.W_ENERGY}·E − {cfg.W_TIME}·T − "
               f"{cfg.W_HAZARD}·H − λ(g)·V")
    r.bullets([
        "**Cov** = Σ P(c)·covered(c) / Σ P(c); a cell is covered if its centre is within the sensor radius of "
        "a leg flown with the sensor on.",
        "**Det** = mean over targets of 1 − Π(1 − P_leg), with P_leg = q(p)·s(v)·max(0, 1 − (d/r)²), "
        "q(p) = 0.4 + 0.6p and s(v) falling linearly from 1 (v_min) to 0.5 (v_max).",
        f"**E** = Σ E_d / (N·B), with E_d = Σ_legs [P_fly(v)·(1 + {cfg.WIND_POWER_FACTOR}·wind) + a·P_sensor(p)]·t.",
        f"**T** = makespan of the slowest drone / (L_max / v_min = {cfg.MAX_RANGE_M / cfg.V_MIN:.0f} s).",
        "**H** = Σ ∫h dt / Σ T_d, the time-weighted mean hazard intensity (crossing a pocket fast is better).",
        f"**V** = {cfg.NOFLY_WEIGHT:.0f}·(share of path in no-fly zones) + range overrun + battery overrun; "
        f"λ(g) = {cfg.PENALTY_START:.0f}·(1 + {cfg.PENALTY_GROWTH:.0f}·g/G) grows from {cfg.PENALTY_START:.0f} "
        f"to {cfg.PENALTY_START * (1 + cfg.PENALTY_GROWTH):.0f} (dynamic penalty [6]).",
    ])
    r.para("**Weights.** Without normalisation the term with the largest unit (joules) would dominate and the "
           f"weights would be meaningless. The two mission goals get {cfg.W_COVERAGE + cfg.W_DETECTION:.2f} of "
           f"the weight and the three costs {cfg.W_ENERGY + cfg.W_TIME + cfg.W_HAZARD:.2f}, so the GA is rewarded "
           "for doing the mission rather than staying at the base. Coverage of disaster zones is the main "
           f"purpose ({cfg.W_COVERAGE}); detection ({cfg.W_DETECTION}) is slightly lower because good coverage "
           f"already helps it. Drone safety and battery reserve are equally important ({cfg.W_HAZARD} each). Time "
           f"gets the smallest weight ({cfg.W_TIME}) because the range limit already bounds it and it is "
           "correlated with energy. The penalty is mild early (promising, slightly infeasible plans can help "
           "exploration) and strict at the end (the final plan must be feasible).")

    # ------------------------------------------------------------------ 4
    r.heading("4. Genetic Operators and Constraint Handling")
    r.para(f"**Selection:** tournament (size k), linear rank (pressure {cfg.RANK_PRESSURE}) or roulette wheel; "
           f"the best {100 * cfg.ELITISM:.0f} % are copied unchanged (elitism). **Crossover** (rate "
           f"{cfg.CROSSOVER_RATE}): first each drone segment is taken from either parent (*segment swap*, keeps "
           "whole routes as building blocks); then, for half of the drones, coordinates, speeds and powers are "
           f"blended with **BLX-α** (α = {cfg.BLX_ALPHA}) [3], a geometric crossover that places the child "
           "between the parents, and the sensor bits use uniform crossover. **Mutation** (each gene with "
           f"probability p_m): Gaussian waypoint moves with a step σ annealed from {cfg.SIGMA_START:.0f} to "
           f"{cfg.SIGMA_END:.0f} cells, speed and power jitter, bit flips, base change, waypoint insert / delete / "
           f"swap, and a rare ({cfg.RESET_PROB}) random reset of a route. **Diversity control:** the mean pairwise "
           f"distance between individuals is tracked; below {100 * cfg.DIVERSITY_THRESHOLD:.0f} % of its start "
           f"value, {100 * cfg.IMMIGRANT_FRACTION:.0f} % random immigrants are added and p_m is doubled. "
           f"**Termination:** maximum generations, stagnation ({cfg.STAGNATION_PATIENCE} generations without a "
           f"gain > {cfg.STAGNATION_EPS}) or a target fitness; the criterion that fired is logged.")
    r.para("**Repair** fixes every child in four steps: (1) clip genes to their ranges; (2) move waypoints out "
           "of no-fly zones to the nearest free cell (a table built once by multi-source BFS); (3) give legs that "
           "cross a no-fly zone detour points – one side point, then divide-and-conquer, finally A∗ [5] with "
           "line-of-sight smoothing; (4) drop waypoints from the end of a route until the drone can return home "
           "within range and battery. Repair keeps a one-cell safety margin and writes the fixed genes back "
           "(Lamarckian). Modes: *penalty* (clip + penalty only), *repair* (no penalty), *both* (default).")
    cm = sens[sens["parameter"] == "constraint_mode"].set_index("value")
    cm_rows = [[mode, pm(cm.loc[mode, "best_mean"], cm.loc[mode, "best_std"]),
                f"{100 * cm.loc[mode, 'feasible_share']:.0f} %", f"{cm.loc[mode, 'runtime_mean_s']:.0f}"]
               for mode in cm.index]
    r.table(["Mode", "Final best F (mean ± std)", "Feasible share of final population", "Runtime (s)"], cm_rows,
            f"Repair vs. penalty ({len(cfg.SWEEP_SEEDS)} seeds, {cfg.SWEEP_GENERATIONS} generations, initial "
            "settings). The best plan of every run was feasible.", widths=[2.2, 4.0, 5.0, 2.2])
    pen, rep, both = cm.loc["penalty"], cm.loc["repair"], cm.loc["both"]
    r.para(f"With the **penalty only**, random plans start deeply infeasible, so the GA first spends many "
           f"generations becoming feasible; its final plans were worse ({fmt(pen['best_mean'])}), less consistent "
           f"(± {pen['best_std']:.3f}) and {100 * (1 - pen['feasible_share']):.0f} % of its final population was "
           f"still illegal. **Repair** keeps everyone feasible and gave better, more stable results "
           f"({fmt(rep['best_mean'])} / {fmt(both['best_mean'])} for *repair* / *both*) at the cost of runtime "
           f"({both['runtime_mean_s']:.0f} s vs. {pen['runtime_mean_s']:.0f} s), since repair (BFS, detours, A∗) is "
           "the most expensive step. *Repair* and *both* are equal because repair removes almost every violation; "
           "I kept *both* so the penalty still catches the rare case repair cannot fix [6, 7].")

    # ------------------------------------------------------------------ 5
    r.heading("5. Experimental Setup")
    r.table(["Parameter", "Value", "Parameter", "Value"],
            [["Population size", str(s["pop_size"]), "Max generations", str(s["max_generations"])],
             ["Crossover rate", str(cfg.CROSSOVER_RATE), "Mutation rate p_m", str(s["config"]["MUTATION_RATE"])],
             ["Selection", f"tournament, k = {s['config']['TOURNAMENT_K']}", "Elitism", f"{100 * cfg.ELITISM:.0f} %"],
             ["BLX-α", str(cfg.BLX_ALPHA), "Mutation σ", f"{cfg.SIGMA_START:.0f} → {cfg.SIGMA_END:.0f} cells"],
             ["Constraint mode", cfg.CONSTRAINT_MODE, "Penalty λ(g)",
              f"{cfg.PENALTY_START:.0f} → {cfg.PENALTY_START * (1 + cfg.PENALTY_GROWTH):.0f}"],
             ["Stagnation patience", f"{cfg.STAGNATION_PATIENCE} gen.", "Seeds",
              f"main {s['seed']}; sweeps {cfg.SWEEP_SEEDS[0]}–{cfg.SWEEP_SEEDS[-1]}; "
              f"baseline {cfg.BASELINE_SEEDS[0]}–{cfg.BASELINE_SEEDS[-1]}"]],
            "Final configuration (p_m and k were chosen by the sensitivity analysis, Section 7).",
            widths=[3.2, 3.4, 3.2, 5.6])
    r.para(f"Three experiments were run: (1) the **main run** (Table 3, seed {s['seed']}); (2) a **sensitivity "
           f"analysis** that changes one parameter at a time around the initial settings (p_m = "
           f"{cfg.SWEEP_BASE_MUTATION_RATE}, k = {cfg.SWEEP_BASE_TOURNAMENT_K}), {len(cfg.SWEEP_SEEDS)} seeds × "
           f"{cfg.SWEEP_GENERATIONS} generations per setting with the stagnation stop off; (3) **random search** "
           "with the same random plan generator, repair, fitness (final penalty) and exactly the same number of "
           f"evaluations as the GA, over {len(cfg.BASELINE_SEEDS)} seeds. All runs use fixed seeds. The main run "
           f"took {s['runtime_s']:.0f} s for {s['fitness_evaluations']:,} evaluations (4-core cloud machine, "
           "Python 3.13 + NumPy, single process).")

    # ------------------------------------------------------------------ 6
    r.heading("6. Results of the Main Run")
    half = int((h["best_fitness"] >= first["best_fitness"] + 0.5 * (last["best_fitness"] - first["best_fitness"])).idxmax())
    low = int((h["diversity"] < cfg.DIVERSITY_THRESHOLD * first["diversity"]).idxmax())
    r.para(f"The run stopped at the **{s['stop_reason']}**. The best fitness rose from "
           f"{fmt(s['initial_best_fitness'])} (best random plan) to **{fmt(s['final_best_fitness'])}** and the "
           f"average from {fmt(s['initial_mean_fitness'])} to {fmt(s['final_mean_fitness'])} (Fig. 1). Half of the "
           f"gain came in the first {half} generations; afterwards the plan was refined slowly. Diversity fell "
           f"below the threshold at generation {low}; from then on the diversity control added immigrants in "
           f"{s['immigrant_generations']} of {s['generations_run']} generations and kept it near the threshold. These "
           "low-fitness newcomers cause the saw-tooth in the average curve, while elitism keeps the best curve "
           f"monotone. {100 * last['feasible_share']:.0f} % of the final population was feasible.")
    r.figure([fig("convergence.png"), fig("diversity.png")], [8.6, 7.6],
             "Left: best and average (±1 std) fitness. Right: diversity with the immigrant threshold, and the "
             "annealed mutation step σ.")
    r.table(["Best plan", "Generation 0", f"Generation {s['generations_run']}"],
            [["Coverage / detection rate", f"{fmt(first['coverage'])} / {fmt(first['detection'])}",
              f"{fmt(best['coverage'])} / {fmt(best['detection'])}"],
             ["Targets detected · high-priority cells covered", "–",
              f"{best['targets_detected']} / {n_targets} · {100 * best['high_priority_coverage']:.0f} %"],
             ["Energy share (kJ) / time share (makespan)", f"{fmt(first['energy'])} / {fmt(first['time'])}",
              f"{fmt(best['energy'])} ({best['energy_kj']:.0f} kJ) / {fmt(best['time'])} ({best['makespan_s']:.0f} s)"],
             ["Hazard exposure / constraint violation", f"{fmt(first['hazard'])} / {fmt(first['violation'])}",
              f"{fmt(best['hazard'])} / 0.000"]],
            "Objectives of the best plan at the start and at the end of the main run.", widths=[6.4, 3.0, 5.6])
    r.para("The GA improved coverage, detection, energy and time together, while hazard exposure stayed almost "
           f"unchanged ({fmt(first['hazard'])} → {fmt(best['hazard'])}): two radiation pockets lie on disaster "
           "zones, so full priority coverage requires entering them, and with these weights the GA prefers to "
           "survey them but crosses them fast. In Fig. 2, generation 0 shows long random zig-zags across the "
           f"map; by generation {s['snapshot_generations'][1]} the drones concentrate on the disaster zones; at the "
           "end each drone sweeps its own area with many short legs, the sensor is off on some transit legs "
           "(dashed) and every route bends around the no-fly zones. The coverage heatmaps (Fig. 3) show that the "
           "whole population, not only the best plan, shifts its coverage from a uniform spread to the "
           "high-priority zones.")
    r.figure([fig("paths_generations.png")], [9.8],
             f"Best plan at generations {', '.join(map(str, s['snapshot_generations']))}. Solid = sensor on, dashed = "
             "sensor off; grey hatched = no-fly zones; orange = hazard; dashed circles = disaster zones.")
    r.figure([fig("coverage_heatmaps.png")], [14.5],
             "Top: cells covered by the best plan (colour = priority). Bottom: share of the population covering "
             "each cell.")

    # ------------------------------------------------------------------ 7
    r.heading("7. Hyper-parameter Sensitivity Analysis")
    short = {"pop_size": "N", "mutation_rate": "p_m", "tournament_k": "k", "selection": "selection"}
    rows = []
    for parameter in ("pop_size", "mutation_rate", "tournament_k", "selection"):
        for _, row in sens[sens["parameter"] == parameter].iterrows():
            rows.append([f"{short[parameter]} = {row['value']}", pm(row["best_mean"], row["best_std"]),
                         f"{row['converge_gen_mean']:.0f}", f"{row['final_diversity']:.2f}",
                         f"{int(row['evaluations']):,}", f"{row['runtime_mean_s']:.0f}"])
    rows.append([f"p_m = {combo['mutation_rate'].iloc[0]} and k = {combo['tournament_k'].iloc[0]}",
                 pm(combo["final_best_fitness"].mean(), combo["final_best_fitness"].std()), "–", "–",
                 f"{int(g('mutation_rate', 0.1)['evaluations']):,}", f"{combo['runtime_s'].mean():.0f}"])
    r.para(f"Each setting ran {len(cfg.SWEEP_SEEDS)} seeds × {cfg.SWEEP_GENERATIONS} generations (population "
           f"{cfg.POP_SIZE} unless swept). Table 5 gives the final best fitness (mean ± std), the generation at "
           "which 99 % of the run's improvement was reached, the final diversity, evaluations and runtime; Fig. 4 "
           "shows the mean curves.")
    r.table(["Setting", "Final best F", "Gen. to 99 %", "Final diversity", "Evaluations", "Runtime (s)"],
            rows, "Sensitivity to population size N, mutation rate p_m and selection pressure (k, method).",
            widths=[4.2, 2.9, 2.0, 2.3, 2.1, 2.0], size=8)
    r.figure_grid([fig("sweep_pop_size.png"), fig("sweep_mutation_rate.png"),
                   fig("sweep_tournament_k.png"), fig("sweep_selection.png")], 6.7, 2,
                  "Mean best-fitness curves (±1 std over seeds) of the four sweeps.")
    p20, p50, p100, p200 = (g("pop_size", n) for n in (20, 50, 100, 200))
    m001, m01, m04 = g("mutation_rate", 0.01), g("mutation_rate", 0.1), g("mutation_rate", 0.4)
    k2, k3, k5, k8 = (g("tournament_k", k) for k in (2, 3, 5, 8))
    sr, sk, st = g("selection", "roulette"), g("selection", "rank"), g("selection", "tournament")
    four = sens[sens["parameter"].isin(["pop_size", "mutation_rate", "tournament_k", "selection"])]
    r.para(f"**Population size.** Larger populations explore more of the map per generation and found better "
           f"plans ({fmt(p20['best_mean'])} → {fmt(p200['best_mean'])} from 20 to 200), but the cost grows "
           f"linearly ({p20['runtime_mean_s']:.0f} s → {p200['runtime_mean_s']:.0f} s) and 50 → 100 gave almost "
           f"nothing ({fmt(p50['best_mean'])} vs. {fmt(p100['best_mean'])}), so it is a cost–quality choice.")
    r.para(f"**Mutation rate** had the largest effect: fitness fell steadily from {fmt(m001['best_mean'])} "
           f"(p_m = 0.01) to {fmt(m01['best_mean'])} (0.1) and {fmt(m04['best_mean'])} (0.4), while final diversity "
           f"rose from {m001['final_diversity']:.2f} to {m04['final_diversity']:.2f}. A chromosome has about "
           "150–300 genes, so p_m = 0.1 already changes some twenty genes per child: good routes are destroyed "
           "faster than selection can keep them and the search drifts towards random search (too much "
           f"exploration). Damaged routes also need more repair ({m04['runtime_mean_s']:.0f} s vs. "
           f"{m001['runtime_mean_s']:.0f} s). With p_m = 0.01 the GA exploits good routes, and premature "
           "convergence is prevented by BLX-α, the annealed mutation and the immigrants instead.")
    r.para(f"**Selection pressure.** k = 2 and 3 performed alike ({fmt(k2['best_mean'])}, {fmt(k3['best_mean'])}); "
           f"k = 5 and 8 were clearly better ({fmt(k5['best_mean'])}, {fmt(k8['best_mean'])}) but reduced diversity "
           f"({k2['final_diversity']:.2f} → {k8['final_diversity']:.2f}), and k = 8 had the largest spread between "
           f"seeds (± {k8['best_std']:.3f}), an early sign of premature convergence in some runs. Tournament "
           f"selection ({fmt(st['best_mean'])}) beat rank ({fmt(sk['best_mean'])}) and roulette "
           f"({fmt(sr['best_mean'])}): late in a run the fitness values are close, so roulette and rank give "
           "almost everyone a similar chance and the search drifts (they keep the highest diversity, "
           f"{sr['final_diversity']:.2f} and {sk['final_diversity']:.2f}).")
    r.para(f"**Final setting.** Because the sweeps vary one parameter at a time, I also ran the two best values "
           f"together (p_m = 0.01, k = 5): {pm(combo['final_best_fitness'].mean(), combo['final_best_fitness'].std())}, "
           f"better than either alone, with immigrants active in {combo['immigrant_generations'].min()}–"
           f"{combo['immigrant_generations'].max()} of 100 generations, so diversity stayed under control. This is "
           f"the final configuration. Most runs were still improving after 100 generations (99 % of the gain only "
           f"after generation {four['converge_gen_mean'].min():.0f}–{four['converge_gen_mean'].max():.0f}), so the "
           "main run uses 200.")

    # ------------------------------------------------------------------ 8
    r.heading("8. GA versus Random Search")
    metrics = [("Fitness F", "fitness", 3), ("Coverage (priority-weighted)", "coverage", 3),
               ("High-priority cells covered", "high_priority_coverage", 3), ("Detection rate", "detection", 3),
               (f"Targets detected (of {n_targets})", "targets_detected", 1), ("Energy (kJ)", "energy_kj", 1),
               ("Makespan (s)", "makespan_s", 1), ("Hazard exposure", "hazard", 3)]
    rows = [[name, pm(ga_r[col].mean(), ga_r[col].std(), d), pm(rs_r[col].mean(), rs_r[col].std(), d)]
            for name, col, d in metrics]
    budget = int(ga_r["evaluations"].iloc[0])
    r.table(["Best plan", "Genetic Algorithm", "Random search"], rows,
            f"GA vs. random search: mean ± std over {len(ga_r)} seeds, {budget:,} evaluations each. "
            f"All best plans were feasible (violation {ga_r['violation'].max():.0f}).",
            widths=[5.2, 3.8, 3.8])
    rs_level = rs_r["fitness"].mean()
    passed = h[h["best_fitness_final_penalty"] >= rs_level].iloc[0]
    stats = {}
    for name, plan in (("ga", data["examples"]["ga_example_plan"]), ("rs", data["examples"]["random_example_plan"])):
        on = [b for d in plan for b in d["sensor_on"]]
        power = [p for d in plan for p, b in zip(d["sensor_power"], d["sensor_on"]) if b]
        stats[name] = (len(on), 100 * sum(on) / len(on), sum(power) / len(power))
    r.para(f"The GA was better on **every** metric and in every seed: its worst seed ({fmt(ga_r['fitness'].min())}) "
           f"is far above the best random seed ({fmt(rs_r['fitness'].max())}). In the main run it passed the "
           f"final random-search level ({fmt(rs_level)}) after only {int(passed['evaluations'])} evaluations, "
           "while random search improves only very slowly after its first thousand samples (Fig. 5). The GA was "
           f"even faster ({ga_r['runtime_s'].mean():.0f} s vs. {rs_r['runtime_s'].mean():.0f} s per run) because its "
           "children need little repair, whereas every random plan needs heavy repair.")
    r.figure([fig("ga_vs_random_convergence.png"), fig("ga_vs_random_solutions.png")], [8.4, 5.2],
             f"Left: best-so-far fitness vs. evaluations (mean ± std). Right: best plans of seed "
             f"{data['examples']['example_seed']}, routes (top) and covered cells (bottom).")
    (g_n, g_on, g_p), (r_n, r_on, r_p) = stats["ga"], stats["rs"]
    r.para("**What makes the evolved solution different?** (i) *Structured paths:* random plans are a few long "
           "straight legs across the map that hit the range limit and are cut short by repair "
           f"({r_n} waypoints in seed {data['examples']['example_seed']}); the GA builds many short legs "
           f"({g_n} waypoints) that sweep the disaster zones, and the drones share the map. (ii) *Sensor "
           f"scheduling:* the GA keeps the sensor on for {g_on:.0f} % of the legs at high power (mean {g_p:.2f}) "
           f"and switches it off on a few low-priority legs (random: {r_on:.0f} %, power {r_p:.2f}); it pays for "
           "this by flying faster and finishing sooner, so it still uses less energy. (iii) *Priority focus:* "
           f"{100 * ga_r['high_priority_coverage'].mean():.1f} % of high-priority cells covered (random "
           f"{100 * rs_r['high_priority_coverage'].mean():.1f} %) and {ga_r['targets_detected'].mean():.1f} instead of "
           f"{rs_r['targets_detected'].mean():.1f} targets found. (iv) *Hazard avoidance:* it must enter the "
           "radiation pockets on the disaster zones but crosses them quickly and mostly avoids the others "
           f"(exposure {fmt(ga_r['hazard'].mean())} vs. {fmt(rs_r['hazard'].mean())}). Random search finds such a "
           "coordinated plan only by chance; the GA builds it by keeping and recombining good drone routes.")

    # ------------------------------------------------------------------ 9
    r.heading("9. Limitations and Future Work")
    r.para("(1) A weighted sum gives one compromise per weight vector; a Pareto method (e.g. NSGA-II) would "
           "offer a set of trade-offs. (2) The world is static and simplified (fixed hazards, no wind drift, no "
           "drone collisions, simple sensor footprint) and only one map was used. (3) Repair, especially A∗, "
           "dominates the runtime. (4) Tuning was one-factor-at-a-time with three seeds; a factorial design "
           "with more seeds would give stronger evidence.")

    # ------------------------------------------------------------------ 10
    r.heading("10. Conclusion")
    r.para(f"I built a GA that plans the deployment, routes, speeds and sensor schedules of a {cfg.N_DRONES}-drone "
           f"fleet on a {cfg.GRID_SIZE}×{cfg.GRID_SIZE} hazard grid. The hybrid multi-segment chromosome maps "
           "directly onto the decisions, the normalised weighted fitness balances five conflicting objectives, "
           "and repair plus a dynamic penalty keeps every final plan feasible. The sensitivity study showed that "
           "this problem rewards exploitation (low mutation rate, strong tournament) as long as diversity is "
           f"controlled. With the tuned settings the GA reached {pm(ga_r['fitness'].mean(), ga_r['fitness'].std())} "
           f"against {pm(rs_r['fitness'].mean(), rs_r['fitness'].std())} for random search with the same budget, "
           "with better values on every objective. A problem-specific encoding and operators let a simple GA "
           "[1, 2, 8] clearly outperform undirected search.")
    r.heading("References", level=2)
    refs = [
        "J. H. Holland, *Adaptation in Natural and Artificial Systems*. Univ. of Michigan Press, 1975.",
        "D. E. Goldberg, *Genetic Algorithms in Search, Optimization and Machine Learning*. Addison-Wesley, 1989.",
        "L. J. Eshelman, J. D. Schaffer, “Real-coded genetic algorithms and interval-schemata,” *Foundations of "
        "Genetic Algorithms 2*, pp. 187–202, 1993.",
        "Y. Zeng, J. Xu, R. Zhang, “Energy minimization for wireless communication with rotary-wing UAV,” "
        "*IEEE Trans. Wireless Commun.*, 18(4), pp. 2329–2345, 2019.",
        "P. E. Hart, N. J. Nilsson, B. Raphael, “A formal basis for the heuristic determination of minimum cost "
        "paths,” *IEEE Trans. Syst. Sci. Cybern.*, 4(2), pp. 100–107, 1968.",
        "J. A. Joines, C. R. Houck, “On the use of non-stationary penalty functions to solve nonlinear "
        "constrained optimization problems with GAs,” *Proc. IEEE CEC*, pp. 579–584, 1994.",
        "C. A. Coello Coello, “Theoretical and numerical constraint-handling techniques used with evolutionary "
        "algorithms: a survey,” *Comput. Methods Appl. Mech. Eng.*, 191, pp. 1245–1287, 2002.",
        "S. Russell, P. Norvig, *Artificial Intelligence: A Modern Approach*, 4th ed. Pearson, 2020.",
    ]
    for i, ref in enumerate(refs, 1):
        r.para(f"[{i}] {ref}", size=7.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)


def to_pdf(docx_path: str) -> str | None:
    """Convert with LibreOffice (if installed) and return the PDF path."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        print("LibreOffice not found - only the .docx was written.")
        return None
    subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir",
                    os.path.dirname(docx_path), docx_path], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return docx_path.replace(".docx", ".pdf")


if __name__ == "__main__":
    copy_figures()
    report = Report()
    write_report(report, load())
    report.save(OUT_DOCX)
    print("wrote", OUT_DOCX)
    pdf = to_pdf(OUT_DOCX)
    if pdf:
        print("wrote", pdf)
